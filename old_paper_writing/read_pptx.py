from pptx import Presentation

prs = Presentation('results.pptx')

for slide_num, slide in enumerate(prs.slides, 1):
    print(f"\n{'='*60}")
    print(f"SLIDE {slide_num}")
    print('='*60)
    
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            text = shape.text.strip()
            if text:
                print(text)
        
        if shape.has_table:
            table = shape.table
            print("\n[TABLE]")
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                print(" | ".join(row_data))
